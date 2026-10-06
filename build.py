import numpy as np
from scipy.stats import norm
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.9))
sec.top_margin = sec.bottom_margin = Inches(0.8)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(4)

def P(text="", bold=False, italic=False, size=None, align=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def RP(parts):
    """paragraph with mixed bold runs: list of (text, bold)"""
    p = doc.add_paragraph()
    for t, b in parts:
        r = p.add_run(t); r.bold = b
    return p

def H(text, lvl=1):
    h = doc.add_heading(text, level=lvl)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return h

def shade(cell, color="DCE6F1"):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:fill"), color)
    tcPr.append(s)

def T(header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def IMG(path, w=5.6, cap=None):
    doc.add_picture(path, width=Inches(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap: P(cap, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

def CODE(text, size=8.5, fill="F4F4F4"):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; shade(c, fill); c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line); r.font.name = "Courier New"; r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        if i < len(lines) - 1: r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

f4 = lambda x: f"{x:.4f}"

# ---------------- header ----------------
P("Kanaka Sarat Siripurapu", bold=True, size=13, after=0)
P("SJSU ID: 019132776", bold=True, size=11, after=0)
P("CMPE 253, AI Threat Intelligence", size=11, after=0)
P("Homework 2", size=11, after=10)
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Homework 2: Fairness, Detection, Robustness and Privacy"); r.bold = True; r.font.size = Pt(15)
P("Rounding: intermediate values are kept to 4 decimal places and final answers are given to 3 significant "
  "figures (or 4 decimals where a table needs them). All plots in Part A were made with Matplotlib from the "
  "numbers I worked out by hand. The Part B code uses only NumPy, SciPy and Matplotlib.", italic=True, size=9.5)

# =============== PART A ===============
H("Part A: Numerical Questions", 1)

# ---------- NQ1 ----------
H("NQ1. Fairness metrics for groups C and D", 2)
P("(a) Per-group metrics", bold=True)
P("Group C: TP = 72, FP = 18, TN = 140, FN = 30")
P("N = 72 + 18 + 140 + 30 = 260\n"
  "Selection rate = (TP + FP) / N = 90 / 260 = 0.3462\n"
  "TPR = TP / (TP + FN) = 72 / 102 = 0.7059\n"
  "FPR = FP / (FP + TN) = 18 / 158 = 0.1139\n"
  "Precision = TP / (TP + FP) = 72 / 90 = 0.8000")
P("Group D: TP = 50, FP = 40, TN = 90, FN = 20")
P("N = 50 + 40 + 90 + 20 = 200\n"
  "Selection rate = 90 / 200 = 0.4500\n"
  "TPR = 50 / 70 = 0.7143\n"
  "FPR = 40 / 130 = 0.3077\n"
  "Precision = 50 / 90 = 0.5556")
P("(b) Summary table", bold=True)
T(["Metric", "Group C", "Group D", "Gap (D - C)"], [
    ["N", "260", "200", "-60"],
    ["Selection rate", "0.3462", "0.4500", "+0.1038"],
    ["TPR", "0.7059", "0.7143", "+0.0084"],
    ["FPR", "0.1139", "0.3077", "+0.1938"],
    ["Precision", "0.8000", "0.5556", "-0.2444"],
])
P("(c) Which criterion holds and which breaks", bold=True)
P("Equal opportunity is the closest to being met. The TPR gap is only 0.0084, so qualified applicants in C and D "
  "get picked at almost the same rate. Predictive parity is broken the worst: precision drops from 0.80 in C to "
  "0.56 in D, a gap of 0.244, which is the largest gap in absolute size. In plain terms, a \"yes\" from the model "
  "is much less reliable for group D. Equalized odds also fails, because it needs FPR to match as well as TPR, and "
  "the FPR gap is 0.194 (group D's unqualified applicants get through almost three times as often). Demographic "
  "parity is off by about 10 points, so it fails as well, but not as badly as these two.")

# ---------- NQ2 ----------
H("NQ2. Are the NQ1 gaps statistically significant?", 2)
P("(a) Denominators", bold=True)
T(["Metric", "Denominator", "n (C)", "n (D)"], [
    ["Selection rate", "N", "260", "200"],
    ["TPR", "TP + FN", "102", "70"],
    ["FPR", "FP + TN", "158", "130"],
    ["Precision", "TP + FP", "90", "90"],
])
P("(b) SE, z and p for each gap", bold=True)
C = dict(sel=(90/260, 260), tpr=(72/102, 102), fpr=(18/158, 158), prec=(72/90, 90))
D = dict(sel=(90/200, 200), tpr=(50/70, 70), fpr=(40/130, 130), prec=(50/90, 90))
lab = dict(sel="Selection rate", tpr="TPR", fpr="FPR", prec="Precision")
rows = []
for m in C:
    p1, n1 = C[m]; p2, n2 = D[m]
    v1 = p1*(1-p1)/n1; v2 = p2*(1-p2)/n2; se = np.sqrt(v1+v2); gap = p2-p1; z = gap/se; pv = 2*norm.sf(abs(z))
    P(f"{lab[m]}: p1(1-p1)/n1 = {p1:.4f}({1-p1:.4f})/{n1} = {v1:.6f};  p2(1-p2)/n2 = {p2:.4f}({1-p2:.4f})/{n2} = {v2:.6f}\n"
      f"    SE = sqrt({v1:.6f} + {v2:.6f}) = sqrt({v1+v2:.6f}) = {se:.4f};   z = {gap:+.4f} / {se:.4f} = {z:+.3f};   "
      f"p = 2(1 - Phi(|z|)) = {pv:.3g}", size=10)
    rows.append([lab[m], f"{gap:+.4f}", f"{se:.4f}", f"{z:+.3f}", f"{pv:.3g}", f"[{gap-1.96*se:+.3f}, {gap+1.96*se:+.3f}]", "Y" if pv < 0.05 else "N"])
T(["Metric", "Gap", "SE", "z", "p-value", "95% CI", "Sig. at 0.05?"], rows)
P("(c) Conclusion", bold=True)
P("The selection-rate, FPR and precision gaps are all significant at alpha = 0.05. FPR and precision are far past "
  "the cutoff (p below 0.001). The TPR gap is not significant at all (p = 0.905). Its interval goes from about "
  "-0.13 to +0.15, so it is basically noise. That fits what I said in NQ1(c).")
IMG("plots/nq2_ci.png", 5.2, "Figure NQ2. Each gap with gap +/- 1.96 x SE. Only the TPR interval crosses zero.")

# ---------- NQ3 ----------
H("NQ3. Impossibility result with base rates 0.45 and 0.30", 2)
P("(a) Operating point 1: TPR = 0.65, FPR = 0.20", bold=True)
P("Group E (pi = 0.45): numerator = 0.45(0.65) = 0.2925; (1 - 0.45)(0.20) = 0.1100\n"
  "    PPV_E = 0.2925 / (0.2925 + 0.1100) = 0.2925 / 0.4025 = 0.7267\n"
  "Group F (pi = 0.30): numerator = 0.30(0.65) = 0.1950; (0.70)(0.20) = 0.1400\n"
  "    PPV_F = 0.1950 / 0.3350 = 0.5821\n"
  "Precision gap = 0.7267 - 0.5821 = 0.1446")
P("(b) Operating point 2: TPR = 0.55, FPR = 0.15", bold=True)
P("Group E: 0.45(0.55) = 0.2475; 0.55(0.15) = 0.0825;  PPV_E = 0.2475 / 0.3300 = 0.7500\n"
  "Group F: 0.30(0.55) = 0.1650; 0.70(0.15) = 0.1050;  PPV_F = 0.1650 / 0.2700 = 0.6111\n"
  "Precision gap = 0.7500 - 0.6111 = 0.1389")
T(["Operating point", "TPR", "FPR", "PPV (E)", "PPV (F)", "Gap (E - F)"], [
    ["1", "0.65", "0.20", "0.7267", "0.5821", "0.1446"],
    ["2", "0.55", "0.15", "0.7500", "0.6111", "0.1389"],
])
P("(c) Why the gap never closes", bold=True)
P("Dividing the top and bottom of the PPV formula by pi*TPR gives PPV = 1 / (1 + [(1 - pi)/pi] x [FPR/TPR]). With "
  "equalized odds, both groups share the same FPR/TPR, so the only thing left that differs is (1 - pi)/pi. That is "
  "1.222 for E and 2.333 for F. For the PPVs to be equal those two would have to be equal, which only happens when "
  "the base rates are equal. Moving the operating point from 1 to 2 changed FPR/TPR from 0.308 to 0.273 and the gap "
  "only shrank from 0.145 to 0.139. Pushing FPR toward zero would shrink it more, but it only hits zero at FPR = 0, "
  "which means a perfect classifier. So with real classifiers and pi_E different from pi_F, equalized odds and "
  "predictive parity can't both hold.")

# ---------- NQ4 ----------
H("NQ4. Reweighting by hand", 2)
P("(a) Marginal probabilities (N = 450)", bold=True)
P("P(E) = 250/450 = 0.5556     P(F) = 200/450 = 0.4444\n"
  "P(approved) = 160/450 = 0.3556     P(rejected) = 290/450 = 0.6444")
P("(b) Expected counts and weights", bold=True)
P("Expected count = P(group) x P(label) x 450. The weight is expected / observed.\n"
  "E, approved: 0.5556 x 0.3556 x 450 = 88.889;  weight = 88.889 / 90 = 0.9877\n"
  "E, rejected: 0.5556 x 0.6444 x 450 = 161.111;  weight = 161.111 / 160 = 1.0069\n"
  "F, approved: 0.4444 x 0.3556 x 450 = 71.111;  weight = 71.111 / 70 = 1.0159\n"
  "F, rejected: 0.4444 x 0.6444 x 450 = 128.889;  weight = 128.889 / 130 = 0.9915")
T(["Group x label", "Observed", "Expected", "Weight"], [
    ["E, approved", "90", "88.889", "0.9877"],
    ["E, rejected", "160", "161.111", "1.0069"],
    ["F, approved", "70", "71.111", "1.0159"],
    ["F, rejected", "130", "128.889", "0.9915"],
])
P("(c) Largest and smallest weight", bold=True)
P("F-approved gets the largest weight (1.0159) and E-approved gets the smallest (0.9877). During training each "
  "F-approved example counts a bit more and each E-approved example a bit less, which pulls group and label toward "
  "independence. All four weights are within about 1.6% of 1, though. The approval rates are already close "
  "(36% for E vs 35% for F), so this dataset only needs a small correction.")

# ---------- NQ5 ----------
H("NQ5. Equal-opportunity threshold by bisection", 2)
P("(a) Bisection, 5 iterations", bold=True)
P("TPR_G(t) = 1 - t^1.5 goes down as t goes up. So if TPR at the midpoint is above 0.80, the threshold is too low and "
  "lo moves up to mid. If it is below 0.80, hi moves down to mid.")
P("Iter 1: mid = 0.5;      0.5^1.5 = 0.3536;   TPR = 0.6464 < 0.80, so hi = 0.5\n"
  "Iter 2: mid = 0.25;     0.25^1.5 = 0.1250;  TPR = 0.8750 > 0.80, so lo = 0.25\n"
  "Iter 3: mid = 0.375;    0.375^1.5 = 0.2296; TPR = 0.7704 < 0.80, so hi = 0.375\n"
  "Iter 4: mid = 0.3125;   0.3125^1.5 = 0.1747; TPR = 0.8253 > 0.80, so lo = 0.3125\n"
  "Iter 5: mid = 0.34375;  0.34375^1.5 = 0.2015; TPR = 0.7985 < 0.80, so hi = 0.34375", size=10)
T(["Iter", "lo", "hi", "mid", "TPR(mid)", "Move"], [
    ["1", "0.0000", "1.0000", "0.5000", "0.6464", "hi"],
    ["2", "0.0000", "0.5000", "0.2500", "0.8750", "lo"],
    ["3", "0.2500", "0.5000", "0.3750", "0.7704", "hi"],
    ["4", "0.2500", "0.3750", "0.3125", "0.8253", "lo"],
    ["5", "0.3125", "0.3750", "0.3438", "0.7985", "hi"],
])
P("(b) Final bracket", bold=True)
P("After 5 iterations the bracket is [0.3125, 0.34375]. Its midpoint, 0.3281, is the threshold I would report.")
P("(c) Exact answer", bold=True)
P("1 - t^1.5 = 0.80, so t^1.5 = 0.20 and t = 0.20^(2/3) = exp((2/3) ln 0.2) = exp(-1.0730) = 0.3420.\n"
  "The bisection estimate of 0.3281 is low by 0.0139. That is less than half the bracket width (0.03125), as it "
  "should be. Each extra iteration cuts the error bound in half, so a few more steps would close the gap.")
IMG("plots/nq5_bisect.png", 5.0, "Figure NQ5. TPR_G(t) with the 5 midpoints tried (numbers are the iteration).")

# ---------- NQ6 ----------
H("NQ6. Drift detection with PSI and KS", 2)
P("(a) PSI", bold=True)
e = np.array([.15, .25, .30, .20, .10]); a = np.array([.10, .20, .28, .24, .18])
d = a - e; l = np.log(a / e); c = d * l
rows = [[str(i+1), f"{e[i]:.2f}", f"{a[i]:.2f}", f"{d[i]:+.2f}", f"{a[i]/e[i]:.4f}", f"{l[i]:+.4f}", f"{c[i]:.4f}"] for i in range(5)]
rows.append(["Sum", "1.00", "1.00", "0.00", "", "", f"{c.sum():.4f}"])
T(["Bin", "e", "a", "a - e", "a / e", "ln(a/e)", "(a-e) ln(a/e)"], rows)
P(f"PSI = 0.0203 + 0.0112 + 0.0014 + 0.0073 + 0.0470 = {c.sum():.4f}, so about 0.0871. Most of it comes from bin 5, "
  "where the share went from 10% to 18%.")
P("(b) KS", bold=True)
ce, ca = np.cumsum(e), np.cumsum(a)
T(["Bin", "Cum. expected", "Cum. actual", "|difference|"],
  [[str(i+1), f"{ce[i]:.2f}", f"{ca[i]:.2f}", f"{abs(ca[i]-ce[i]):.2f}"] for i in range(5)])
P("KS = max |difference| = 0.12, which happens at bin 3 (0.70 expected vs 0.58 actual).")
P("(c) Critical value", bold=True)
P("D_crit = 1.36 x sqrt(2/800) = 1.36 x sqrt(0.0025) = 1.36 x 0.05 = 0.068. Since 0.12 > 0.068, the KS test flags drift.")
P("(d) Do they agree?", bold=True)
P("No. PSI (0.087) is under 0.1, so by that rule the model is \"stable\", while KS says the distribution moved. They "
  "can disagree because PSI adds up per-bin terms, and each term stays small when a bin only moves a few points, "
  "but KS looks at the running total, and here every shift is in the same direction (mass moving toward "
  "higher bins), so small changes stack up into a large cumulative gap.")
IMG("plots/nq6_drift.png", 6.2, "Figure NQ6. Left: expected vs actual shares. Right: cumulative curves with the KS gap at bin 3.")

# ---------- NQ7 ----------
H("NQ7. Batch gradient descent for logistic regression", 2)
P("Data: x1 = (1, 1), y1 = 1;  x2 = (-1, 0), y2 = 0;  x3 = (0, -1), y3 = 0.  Start at w = (0, 0), b = 0, eta = 1.")
P("(a) t = 0", bold=True)
P("z_i = w.x_i + b = 0 for all three points, so p_i = sigma(0) = 0.5.\n"
  "J0 = -(1/3)[ln 0.5 + ln(1 - 0.5) + ln(1 - 0.5)] = ln 2 = 0.6931\n"
  "p - y = (-0.5, 0.5, 0.5)\n"
  "grad_w = (1/3)[-0.5(1, 1) + 0.5(-1, 0) + 0.5(0, -1)] = (1/3)(-1, -1) = (-0.3333, -0.3333)\n"
  "grad_b = (1/3)(-0.5 + 0.5 + 0.5) = 0.1667")
P("(b) Updates", bold=True)
P("w1 = (0, 0) - 1(-0.3333, -0.3333) = (0.3333, 0.3333);   b1 = 0 - 0.1667 = -0.1667\n\n"
  "t = 1:  z = (0.3333 + 0.3333 - 0.1667, -0.3333 - 0.1667, -0.3333 - 0.1667) = (0.5, -0.5, -0.5)\n"
  "    p = (sigma(0.5), sigma(-0.5), sigma(-0.5)) = (0.6225, 0.3775, 0.3775)\n"
  "    J1 = -(1/3)[ln 0.6225 + ln 0.6225 + ln 0.6225] = -ln 0.6225 = 0.4741\n"
  "    p - y = (-0.3775, 0.3775, 0.3775)\n"
  "    grad_w = (1/3)(-0.3775 - 0.3775, -0.3775 - 0.3775) = (-0.2517, -0.2517);  grad_b = 0.3775/3 = 0.1258\n"
  "    w2 = (0.3333 + 0.2517, same) = (0.5850, 0.5850);  b2 = -0.1667 - 0.1258 = -0.2925\n\n"
  "t = 2:  z = (1.1700 - 0.2925, -0.5850 - 0.2925, -0.5850 - 0.2925) = (0.8775, -0.8775, -0.8775)\n"
  "    p = (0.7063, 0.2937, 0.2937);  J2 = -ln 0.7063 = 0.3477\n"
  "    grad_w = (1/3)(-0.5874, -0.5874) = (-0.1958, -0.1958);  grad_b = 0.2937/3 = 0.0979\n"
  "    w3 = (0.7808, 0.7808);  b3 = -0.2925 - 0.0979 = -0.3904\n\n"
  "t = 3:  z = (1.1712, -1.1712, -1.1712);  p = (0.7634, 0.2366, 0.2366);  J3 = -ln 0.7634 = 0.2700", size=10)
P("Because of the symmetry in this data set, all three points always end up with the same |z|, so the loss is just "
  "-ln(p1) each time. That made the hand work a lot shorter.", italic=True, size=10)
T(["t", "w1", "w2", "b", "Loss J"], [
    ["0", "0.0000", "0.0000", "0.0000", "0.6931"],
    ["1", "0.3333", "0.3333", "-0.1667", "0.4741"],
    ["2", "0.5850", "0.5850", "-0.2925", "0.3477"],
    ["3", "0.7808", "0.7808", "-0.3904", "0.2700"],
])
P("(c) Decision boundary at t = 2 and the test point", bold=True)
P("0.5850 x1 + 0.5850 x2 - 0.2925 = 0, which simplifies to x1 + x2 = 0.5.\n"
  "For (0.5, 0.5): z = 0.5850(0.5) + 0.5850(0.5) - 0.2925 = 0.2925 > 0, p = sigma(0.2925) = 0.5726. "
  "So the point is classified as 1 (positive). It sits on the same side of the line as x1 = (1, 1).")
IMG("plots/nq7_loss.png", 4.6, "Figure NQ7. Loss J against iteration t.")

# ---------- NQ8 ----------
H("NQ8. Iterative FGSM / PGD by hand", 2)
P("(a) t = 0 and the sign of the gradient", bold=True)
P("f(x0) = 1(0.30) + 2(0.20) - 0.1 = 0.6, p0 = sigma(0.6) = 1 / (1 + e^-0.6) = 1 / 1.5488 = 0.6457.\n"
  "grad_x J = (p0 - 1) w = (-0.3543)(1, 2) = (-0.3543, -0.7086), so sign(grad) = (-1, -1).\n"
  "This sign never changes. Since y = 1 and the sigmoid is always strictly below 1, p - y is always negative, and "
  "w = (1, 2) is fixed with both entries positive. A negative number times a positive vector is always (-, -). So "
  "every step moves both coordinates down by alpha = 0.1.")
P("(b) Four PGD steps (box is x1 in [-0.10, 0.70], x2 in [-0.20, 0.60])", bold=True)
P("t = 1: x = (0.30 - 0.1, 0.20 - 0.1) = (0.20, 0.10), inside the box.  f = 0.20 + 0.20 - 0.1 = 0.30\n"
  "    p = sigma(0.3) = 0.5744,  J = -ln 0.5744 = 0.5544\n"
  "t = 2: x = (0.10, 0.00).  f = 0.10 + 0 - 0.1 = 0.00,  p = 0.5000,  J = ln 2 = 0.6931\n"
  "t = 3: x = (0.00, -0.10).  f = 0 - 0.20 - 0.1 = -0.30,  p = 0.4256,  J = 0.8544\n"
  "t = 4: x = (-0.10, -0.20). Both coordinates are now exactly on the lower edge of the box, so clipping does not "
  "change them.  f = -0.10 - 0.40 - 0.1 = -0.60,  p = 0.3543,  J = 1.0375", size=10)
T(["t", "x_t", "f(x_t)", "p_t", "Loss -ln(p_t)"], [
    ["0", "(0.30, 0.20)", "0.60", "0.6457", "0.4375"],
    ["1", "(0.20, 0.10)", "0.30", "0.5744", "0.5544"],
    ["2", "(0.10, 0.00)", "0.00", "0.5000", "0.6931"],
    ["3", "(0.00, -0.10)", "-0.30", "0.4256", "0.8544"],
    ["4", "(-0.10, -0.20)", "-0.60", "0.3543", "1.0375"],
])
P("(c) Misclassification and the boundary", bold=True)
P("The point is first misclassified at t = 2, where f = 0 (the rule says f <= 0 counts as wrong, and at that point "
  "p = 0.5). Neither coordinate reaches the epsilon box until t = 4, and at t = 4 both hit it at once: "
  "x1 = 0.30 - 0.40 = -0.10 and x2 = 0.20 - 0.40 = -0.20. Any further step would get clipped back, so x would "
  "stay at (-0.10, -0.20).")
IMG("plots/nq8_pgd.png", 4.9, "Figure NQ8. f(x_t) falls by 0.3 per step while the loss rises.")

# ---------- NQ9 ----------
H("NQ9. Gaussian mechanism and privacy accounting", 2)
P("(a) Noise multiplier", bold=True)
P("ln(1.25 / 10^-5) = ln(125,000) = 11.7361;  2 x 11.7361 = 23.4721;  sqrt(23.4721) = 4.8448\n"
  "eps = 2:    sigma = 4.8448 / 2 = 2.42\n"
  "eps = 1:    sigma = 4.8448 / 1 = 4.84\n"
  "eps = 0.5:  sigma = 4.8448 / 0.5 = 9.69")
P("sigma scales as 1/eps, so halving epsilon doubles the noise. A smaller epsilon means a stronger privacy promise, "
  "and the only way to get it is to bury each gradient under more Gaussian noise. More noise means noisier updates, "
  "so the model learns slower and usually ends up less accurate.")
P("(a, second part) Solving for k and the table", bold=True)
P("eps_tight(100) = k sqrt(100) = 10k = 1.5, so k = 0.15.\n"
  "Example row, T = 300: eps_basic = 300 x 0.03 = 9.0; eps_tight = 0.15 x sqrt(300) = 0.15 x 17.3205 = 2.5981; "
  "ratio = 9.0 / 2.5981 = 3.464.")
T(["T", "eps_basic = 0.03T", "eps_tight = 0.15 sqrt(T)", "Ratio basic / tight"], [
    ["100", "3.00", "1.5000", "2.000"],
    ["200", "6.00", "2.1213", "2.828"],
    ["300", "9.00", "2.5981", "3.464"],
    ["400", "12.00", "3.0000", "4.000"],
    ["500", "15.00", "3.3541", "4.472"],
])
P("(b) Overstatement at T = 500", bold=True)
P("15.0 / 3.354 = 4.47, so basic composition overstates the privacy loss by about 4.5x. The ratio is "
  "0.03T / (0.15 sqrt(T)) = 0.2 sqrt(T), so it keeps growing with training length. On the log-log plot the two "
  "lines have slopes 1 and 1/2.")
IMG("plots/nq9_loglog.png", 4.8, "Figure NQ9. eps vs T on log-log axes for both accounting methods.")

# ---------- NQ10 ----------
H("NQ10. Guardrail base rates and a two-stage cascade", 2)
P("(a) Single detector, TPR = 0.90", bold=True)
P("Attacks per day = 2,000,000 x 0.0002 = 400, so true alerts = 400 x 0.90 = 360 (the same for every row).\n"
  "Benign requests = 2,000,000 - 400 = 1,999,600, so false alerts = 1,999,600 x FPR.\n"
  "PPV numerator = 0.0002 x 0.90 = 0.00018; the second term is 0.9998 x FPR.\n"
  "Example, FPR = 1%: PPV = 0.00018 / (0.00018 + 0.009998) = 0.00018 / 0.010178 = 0.0177. "
  "Alerts = 360 + 19,996 = 20,356; at 3 min each that is 61,068 min = 1,017.8 hours.")
T(["FPR", "PPV", "True alerts", "False alerts", "Review hours"], [
    ["5%", "0.00359", "360", "99,980", "5,017.0"],
    ["1%", "0.0177", "360", "19,996", "1,017.8"],
    ["0.5%", "0.0348", "360", "9,998", "517.9"],
    ["0.1%", "0.153", "360", "1,999.6", "118.0"],
])
P("Even at the best FPR, only about 1 alert in 6.5 is a real attack. At 5% FPR, reviewing everything would take "
  "over 5,000 analyst-hours per day, which is not realistic.")
P("(b) Two-stage cascade", bold=True)
P("A benign request has to fool both stages: overall FPR = 0.05 x 0.02 = 0.001 (0.1%).\n"
  "An attack has to be caught by both stages: overall recall = 0.95 x 0.98 = 0.931.")
P("(c) PPV comparison at FPR = 0.1%", bold=True)
P("Single: PPV = 0.0002(0.90) / (0.00018 + 0.9998 x 0.001) = 0.00018 / 0.0011798 = 0.1526\n"
  "Cascade: PPV = 0.0002(0.931) / (0.0001862 + 0.0009998) = 0.0001862 / 0.0011860 = 0.1570\n"
  "The cascade is better, but only a little: PPV goes up by 0.0044 (0.44 percentage points, about 2.9% relative). "
  "The bigger win is recall, 0.931 vs 0.900. That works out to 372 caught attacks per day instead of 360, at the "
  "same false-alarm load.")
IMG("plots/nq10_ppv.png", 6.2, "Figure NQ10. Left: PPV vs FPR on a log x-axis. Right: recall of the single detector vs the cascade.")

# =============== PART B ===============
H("Part B: Coding Questions", 1)
P("Each question below has the full source, the printed output copied from my terminal, and the plot the "
  "script saves. Everything was run with Python 3, NumPy 2.1, SciPy 1.15 and Matplotlib 3.10.", italic=True, size=10)

blocks = [
 ("CQ1. Fairness-metrics function applied to three groups", "cq1", ["plots/cq1_bars.png"],
  "Group P has the median selection rate (0.3393), so the gaps are measured against P. Q selects the most people "
  "but has the worst FPR (+0.18 vs P) and the worst precision (-0.23), so most of its extra selections are false "
  "positives. R has the lowest TPR, about 19 points below P. P looks best on all four metrics, so the "
  "comparison mostly shows how far Q and R fall behind it."),
 ("CQ2. Bootstrap significance test for a metric gap", "cq2", ["plots/cq2_hist.png"],
  "The bootstrap numbers come out close to the z-test from NQ2 (FPR SD 0.0482 vs SE 0.0477; selection-rate SD "
  "0.0460 vs 0.0459), which is a good check on both. The FPR gap stands out more clearly from zero (z about 4.0, "
  "with the 95% interval starting at 0.10) than the selection-rate gap (z about 2.3, interval starting at 0.015). "
  "Sample size alone would point the other way. Selection rate uses all N (260 and 200), while FPR only uses the "
  "negatives (158 and 130), so on its own FPR should be the noisier one. The two SDs end up almost equal, because "
  "FPR values near 0.1 to 0.3 have less binomial variance than selection rates near 0.35 to 0.45. On top of that, "
  "the FPR gap (0.194) is almost twice the selection-rate gap (0.104). So the larger effect is what makes FPR the "
  "clearer result, even though it has fewer samples."),
 ("CQ3. Equal-opportunity threshold search via bisection", "cq3", ["plots/cq3_roc.png"],
  "The matched threshold for G comes out at about 0.501, which is basically H's 0.5. I expected this once I looked "
  "at make_scores: the score distribution for positives is N(0.70, 0.20) in both groups. Only the base rate differs, "
  "and TPR does not depend on the base rate. So the two TPR curves are the same up to sampling noise, and the "
  "search just corrects a 0.002 difference that comes from the random draws. FPR at the matched point is also close "
  "(0.224 vs 0.232) for the same reason. Base rates would start to matter for metrics like precision."),
 ("CQ4. PGD attack, checked against NQ8", "cq4", ["plots/cq4_pgd.png"],
  "Every row matches my hand table from NQ8, so neither the code nor the hand work had a bug. The one thing worth "
  "pointing out is t = 2. In exact math f(x2) = 0, but in floating point 0.1 + 2(0.0) - 0.1 comes out as about "
  "-2.8e-17, because 0.3 - 0.1 - 0.1 is not exactly 0.1 in binary. With the f <= 0 rule the answer is the same "
  "either way. But if the rule were f < 0, float error alone would make the code say t = 2 while the hand answer "
  "says t = 3. I round before printing so the table doesn't show -0.0000."),
 ("CQ5. Sequential Bayesian risk updating", "cq5", ["plots/cq5_posterior.png"],
  "The posterior never gets above 0.50 with this evidence. It ends at 0.2815 after the grid-probing step, which is "
  "the highest it reaches. The reason is the low prior: starting odds are 0.0204, and the product of all four "
  "likelihood ratios is 4 x 1.5 x 0.4 x 8 = 19.2, so the final odds are 0.0204 x 19.2 = 0.392. That is still well "
  "under the 1.0 needed for 50%. Even with only the three positive signals (LR 48 in total), the posterior would "
  "reach about 0.495, just short. The known-good API key is the one piece of evidence that pulls it down, from "
  "0.1091 to 0.0467, a drop of 6.24 percentage points. So a single strong benign signal cancels most of the "
  "early suspicion, and the system would still need one more strong indicator before acting."),
]
for title, name, imgs, disc in blocks:
    H(title, 2)
    P("Source code", bold=True)
    CODE(open(f"{name}.py").read(), size=8)
    P("Printed output", bold=True)
    CODE(open(f"out_{name}.txt").read(), size=8, fill="EEF3EA")
    for im in imgs:
        IMG(im, 5.4 if name != "cq4" else 6.2)
    P("Discussion", bold=True)
    P(disc)

out = "/Users/kanakasarat/Downloads/Kanaka_Sarat_Siripurapu_019132776_CMPE253_HW2.docx"
doc.save(out)
print("saved", out)
